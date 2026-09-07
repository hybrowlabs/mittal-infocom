# Copyright (c) 2026, SAW India and contributors
# For license information, please see license.txt

"""Excel export of the Location Summary report in the layout Tally produces.

Reproduces the letterhead block, the merged four row header per location and the
indented item group tree of the Tally "Location Summary" export, so the file can be
circulated in the format the business already reads.
"""

import json
from io import BytesIO

import frappe
import openpyxl
from frappe import _
from frappe.utils import cint, formatdate
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from mittal_customization.tally.letterhead import get_address_lines

from mittal_customization.mittal_customization.report.location_summary.location_summary import (
	execute,
	get_warehouses,
)

VALUE_REPORT = "Location Summary"
QUANTITY_REPORT = "Location Summary (Quantity)"

QTY_FORMAT = '0" no"'
AMOUNT_FORMAT = "0.00"

PARTICULARS_WIDTH = 89.2

# Tally's own widths are too narrow for figures of this size, so each column is sized
# from the values it actually holds, within these bounds
MIN_QTY_WIDTH = 9
MIN_RATE_WIDTH = 12
MIN_VALUE_WIDTH = 14
MAX_COLUMN_WIDTH = 22
WIDTH_PADDING = 2

THIN_BOTTOM = Border(bottom=Side(style="thin"))


def round_value(value):
	return round(value, 2) if value is not None else None


@frappe.whitelist()
def download_tally_format(filters, report_name=VALUE_REPORT):
	if isinstance(filters, str):
		filters = json.loads(filters)

	filters = frappe._dict(filters)

	if report_name not in (VALUE_REPORT, QUANTITY_REPORT):
		frappe.throw(_("Unknown report {0}").format(report_name))

	# The report the export is asked for decides whether value is included, so its own
	# roles have to be checked here. Permission on the stock ledger alone would let a
	# reader of Location Summary (Quantity) ask for the file that carries the value.
	if not frappe.get_doc("Report", report_name).is_permitted():
		raise frappe.PermissionError

	frappe.has_permission("Stock Ledger Entry", throw=True)

	show_value = report_name == VALUE_REPORT

	_columns, data, _message = execute(filters)

	if not data:
		frappe.throw(_("No stock to report for the selected filters"))

	warehouses = get_warehouses(filters)
	workbook = build_workbook(filters, warehouses, data, report_name, show_value)

	stream = BytesIO()
	workbook.save(stream)

	frappe.response["filename"] = get_file_name(filters, report_name)
	frappe.response["filecontent"] = stream.getvalue()
	frappe.response["type"] = "binary"


def get_file_name(filters, report_name=VALUE_REPORT):
	return "{} {}.xlsx".format(report_name, formatdate(filters.to_date, "dd-MM-yyyy"))


def get_column_groups(filters, warehouses, show_value=True):
	"""The columns shown under each location, plus the total if it is shown.

	Quantity, Rate and Value for Location Summary; Quantity alone for Location Summary
	(Quantity), which is why a group is a list rather than a fixed triple.
	"""

	def columns_for(key):
		columns = [
			{
				"field": f"{key}_qty",
				"label": _("Quantity"),
				"number_format": QTY_FORMAT,
				"formatter": format_qty,
				"min_width": MIN_QTY_WIDTH,
				"bold_on_group": True,
			}
		]

		if show_value:
			columns += [
				{
					"field": f"{key}_rate",
					"label": _("Rate"),
					"number_format": AMOUNT_FORMAT,
					"formatter": format_amount,
					"min_width": MIN_RATE_WIDTH,
					"bold_on_group": False,
				},
				{
					"field": f"{key}_value",
					"label": _("Value"),
					"number_format": AMOUNT_FORMAT,
					"formatter": format_amount,
					"min_width": MIN_VALUE_WIDTH,
					"bold_on_group": True,
				},
			]

		return columns

	groups = [(d.label, columns_for(frappe.scrub(d.name))) for d in warehouses]

	if cint(filters.show_total_column):
		groups.append((_("Total"), columns_for("total")))

	return groups


def last_column_index(groups):
	return 1 + sum(len(columns) for _label, columns in groups)


def build_workbook(filters, warehouses, data, report_name=VALUE_REPORT, show_value=True):
	workbook = openpyxl.Workbook()
	sheet = workbook.active
	sheet.title = report_name

	groups = get_column_groups(filters, warehouses, show_value)

	set_column_widths(sheet, groups, data)
	row = write_letterhead(sheet, filters, groups, report_name)
	row = write_header(sheet, filters, groups, row)
	write_rows(sheet, groups, data, row)

	sheet.freeze_panes = sheet.cell(row=row, column=2).coordinate

	return workbook


def set_column_widths(sheet, groups, data):
	sheet.column_dimensions["A"].width = PARTICULARS_WIDTH

	index = 2
	for _label, columns in groups:
		for column in columns:
			sheet.column_dimensions[get_column_letter(index)].width = measure_width(
				data, column["field"], column["formatter"], column["min_width"]
			)
			index += 1


def format_qty(value):
	return f"{round(value):d} no"


def format_amount(value):
	return f"{value:.2f}"


def measure_width(data, fieldname, formatter, minimum):
	"""Width of the widest figure the column has to show, as Excel renders it."""
	longest = 0

	for row in data:
		value = row.get(fieldname)
		if value is None:
			continue
		longest = max(longest, len(formatter(value)))

	return min(max(longest + WIDTH_PADDING, minimum), MAX_COLUMN_WIDTH)


def merge(sheet, row, start_column, width):
	"""openpyxl treats a single cell range as a merge of nothing, so skip it."""
	if width > 1:
		sheet.merge_cells(
			start_row=row, start_column=start_column, end_row=row, end_column=start_column + width - 1
		)


def write_letterhead(sheet, filters, groups, report_name=VALUE_REPORT):
	"""Company name, address and phone, then the report title and period."""
	last_column = last_column_index(groups)
	lines = [filters.company, *get_address_lines(filters.company)]
	lines.append(_(report_name))
	lines.append(get_period_label(filters))

	for index, line in enumerate(lines):
		row = index + 1
		cell = sheet.cell(row=row, column=1, value=line)
		cell.font = Font(bold=(index == 0))
		cell.alignment = Alignment(horizontal="left")
		merge(sheet, row, 1, last_column)

	return len(lines) + 1


def get_company_caption(filters):
	"""Company name followed by the fiscal year, e.g. "Acme Ltd 26-27"."""
	from erpnext.accounts.utils import get_fiscal_year

	try:
		fiscal_year = get_fiscal_year(filters.to_date, company=filters.company)[0]
	except Exception:
		return filters.company

	return f"{filters.company} {fiscal_year}"


def get_period_label(filters):
	return "{} to {}".format(
		formatdate(filters.from_date, "d-MMM-yy"), formatdate(filters.to_date, "d-MMM-yy")
	)


def write_header(sheet, filters, groups, start_row):
	"""The four merged rows above the columns of each location."""
	period = get_period_label(filters)
	captions = [
		[label for label, _columns in groups],
		[get_company_caption(filters) for _group in groups],
		[period for _group in groups],
		[_("Closing Balance") for _group in groups],
	]

	for offset, values in enumerate(captions):
		row = start_row + offset
		index = 2
		for value, (_label, columns) in zip(values, groups):
			cell = sheet.cell(row=row, column=index, value=value)
			cell.font = Font(bold=(offset == 0))
			cell.alignment = Alignment(horizontal="center")
			merge(sheet, row, index, len(columns))
			index += len(columns)

	# "Particulars" sits against the period row, as Tally places it
	particulars = sheet.cell(row=start_row + 2, column=1, value=_("Particulars"))
	particulars.font = Font(bold=True)

	label_row = start_row + 4
	index = 2
	for _label, columns in groups:
		for column in columns:
			cell = sheet.cell(row=label_row, column=index, value=column["label"])
			cell.font = Font(bold=True)
			cell.alignment = Alignment(horizontal="center")
			cell.border = THIN_BOTTOM
			index += 1

	sheet.cell(row=label_row, column=1).border = THIN_BOTTOM

	return label_row + 1


def write_rows(sheet, groups, data, start_row):
	for offset, source in enumerate(data):
		row = start_row + offset
		is_group = bool(source.get("is_group_row"))

		label = sheet.cell(row=row, column=1, value=source.get("particulars"))
		label.font = Font(bold=is_group)
		# Tally indents each level by two
		label.alignment = Alignment(horizontal="left", indent=(source.get("indent") or 0) * 2)
		if is_group:
			label.border = THIN_BOTTOM

		index = 2
		for _label, columns in groups:
			for column in columns:
				cell = sheet.cell(row=row, column=index, value=round_value(source.get(column["field"])))
				cell.number_format = column["number_format"]
				cell.font = Font(bold=(is_group and column["bold_on_group"]))
				cell.alignment = Alignment(horizontal="right")
				if is_group:
					cell.border = THIN_BOTTOM
				index += 1
